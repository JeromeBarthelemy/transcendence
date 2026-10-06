"""Reading the AssoConnect season export.

Being in the export is what makes someone a member this season. Only the
columns below are read: payment, address, phone, polo size, role and team are
the club's business, not ours, and the file itself is never stored.
"""

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MEMBER_NUMBER = "ID participant"
FIRST_NAME = "Prénom participant"
LAST_NAME = "Nom participant"
BIRTH_DATE = "Date de naissance"
EMAIL = "Adresse email"
GUARDIAN_EMAIL = (
    "Mail du responsable légal à contacter (si différent du mail principal)"
)
IMAGE_RIGHTS_DOCUMENT = (
    "Droit à l'image rempli et signé (formulaire à télécharger sur la première page)"
)
IMAGE_RIGHTS_REFUSAL = (
    "Pour plus de facilité, en cas de refus du droit à l'image, merci de "
    "l'indiquer également ici."
)
DECIDED_ON = "Date de la commande"

EXPECTED_COLUMNS = [
    MEMBER_NUMBER,
    FIRST_NAME,
    LAST_NAME,
    BIRTH_DATE,
    EMAIL,
    GUARDIAN_EMAIL,
    IMAGE_RIGHTS_DOCUMENT,
    IMAGE_RIGHTS_REFUSAL,
    DECIDED_ON,
]


class ExportError(Exception):
    """Read by the club admin, who then fixes the file and imports again."""


@dataclass(frozen=True)
class ExportRow:
    """One member of the season, as the export states them."""

    assoconnect_id: str
    first_name: str
    last_name: str
    birth_date: date
    email: str
    guardian_email: str
    # Only their presence decides: the signed proof stays in AssoConnect
    # and the link is never followed, stored or shown.
    image_rights_signed: bool
    image_rights_refused: bool
    decided_on: date


def read_export(path: Path | str) -> list[ExportRow]:
    """The members of the season, in the order the file states them.

    All or nothing: an export is machine-generated, so one anomaly means the
    wrong file or a changed export, not a row to skip.
    """
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile) as unreadable:
        # A .csv, or a .csv renamed .xlsx: the announced format and the real
        # one are the same problem for whoever reads the message.
        raise ExportError(
            f"Le fichier n'est pas un classeur .xlsx lisible : {path}"
        ) from unreadable
    try:
        # The first sheet, never the active one: the active one is whichever
        # tab was selected when the file was last saved.
        sheet = workbook.worksheets[0]
        # AssoConnect declares the sheet one row long whatever it holds, and
        # read-only openpyxl would trust it and stop after the header.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ExportError("Le fichier est vide : pas même une ligne d'en-tête.")
        position = _positions(header_row)
        # The header is line 1, so the first member is line 2, as the admin
        # sees it in the spreadsheet.
        members = [
            _read_row(row, position, number) for number, row in enumerate(rows, 2)
        ]
    finally:
        workbook.close()
    if not members:
        raise ExportError("Le fichier ne contient aucun membre : rien à importer.")
    _refuse_duplicates(members)
    return members


def _positions(header_row: tuple[Any, ...]) -> dict[str, int]:
    """Where each column sits, read from the header rather than fixed: the
    column order belongs to the club admin, who adds and moves columns."""
    position = {
        header: index for index, header in enumerate(_text(cell) for cell in header_row)
    }
    missing = [column for column in EXPECTED_COLUMNS if column not in position]
    if missing:
        raise ExportError(
            "Colonnes absentes de l'export : " + ", ".join(f"« {c} »" for c in missing)
        )
    return position


def _refuse_duplicates(members: list[ExportRow]) -> None:
    """The member number is what accounts get tied to, so two rows sharing one
    leave the import unable to decide."""
    seen: dict[str, int] = {}
    for number, member in enumerate(members, 2):
        first = seen.get(member.assoconnect_id)
        if first is not None:
            raise ExportError(
                f"Le numéro d'adhérent {member.assoconnect_id} apparaît "
                f"lignes {first} et {number}."
            )
        seen[member.assoconnect_id] = number


def _read_row(row: tuple[Any, ...], position: dict[str, int], number: int) -> ExportRow:
    """One member, `number` being the line the admin sees in the spreadsheet."""

    def value(column: str) -> str:
        return _text(row[position[column]])

    if not value(MEMBER_NUMBER):
        raise ExportError(f"Ligne {number} : numéro d'adhérent absent.")
    return ExportRow(
        assoconnect_id=value(MEMBER_NUMBER),
        first_name=value(FIRST_NAME),
        last_name=value(LAST_NAME),
        birth_date=_as_date(row[position[BIRTH_DATE]], number, BIRTH_DATE),
        email=value(EMAIL),
        guardian_email=value(GUARDIAN_EMAIL),
        image_rights_signed=bool(value(IMAGE_RIGHTS_DOCUMENT)),
        image_rights_refused=bool(value(IMAGE_RIGHTS_REFUSAL)),
        decided_on=_as_date(row[position[DECIDED_ON]], number, DECIDED_ON),
    )


def _text(value: Any) -> str:
    """Everything the export states is text: an empty cell comes back as None
    and a member number as an int."""
    return "" if value is None else str(value).strip()


def _as_date(value: Any, number: int, column: str) -> date:
    """A date column comes back as a real date when the export holds a date
    cell, and as 08/06/1976 when it holds text: the club's file mixes both."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        day, month, year = _text(value).split("/")
        return date(int(year), int(month), int(day))
    except ValueError:
        raise ExportError(
            f"Ligne {number} : « {column} » ne se lit pas comme une date "
            f"jj/mm/aaaa ({_text(value)!r})."
        ) from None
